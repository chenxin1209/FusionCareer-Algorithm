"""
抽岗前关键词预筛：原文未命中新闻学院相关规则则跳过 LLM。

词表来源（可叠加）：
- data/relevance_lexicon.json（本院专业名 + 已确认相关岗位名/媒体词）
- 已抽取 CSV 的岗位名称（仅保留像记者/编辑这类相关名，避免把「后端开发」写进词表）
- 可选：xlsx「岗位汇总」岗位名称列

「产品/市场/行政」等宽词几乎每篇校招都会出现，不能单独作为过筛条件。
纯软件开发 / 人工智能 / 机器人专场若无新闻传播强信号，直接跳过。
抽出的岗位名若是纯技术岗，落盘前再丢一次。
"""
from __future__ import annotations

import argparse
import csv
import json
import os
import re
from typing import Any, Iterable, Optional

from job_structuring import paths

_LEXICON_REL = os.path.join("job_structuring", "data", "relevance_lexicon.json")
_MIN_TITLE_LEN = 2

_TITLE_KEEP = re.compile(
    r"记者|编辑|校对|编导|摄像|摄影|剪辑|主持|主播|出镜|运营|策划|文案|"
    r"媒介|公关|传播|舆情|媒资|通联|评论|融媒体|新媒体|全媒体|美编|编务|"
    r"纪录片|短视频|直播|频道|采访|新闻|"
    r"管培|产品|市场|行政|营销|媒体|职能|广播电视|广告"
)

_STRONG_RE = re.compile(
    r"记者|编导|校对|编务|采编|出镜|融媒体|新媒体|全媒体|广播电视|"
    r"新闻学|广告学|传播学|新闻传播|纪录片|主持人|主播|舆情|媒资|党报|央媒"
)

_TECH_RE = re.compile(
    r"软件开发|软件工程师|人工智能|机器人|算法工程师|机器学习|深度学习|"
    r"嵌入式|芯片研发|后端开发|前端开发|Java开发|C\+\+|golang|Go开发|"
    r"测试开发|运维工程师|数据开发|硬件研发|结构工程师"
)

_KEEP_ARTICLE_RE = re.compile(
    r"管培生|产品经理|记者|编导|融媒体|新媒体|全媒体|宣传岗|市场岗|采编|内容编辑|新媒体编辑"
)

# 几乎每篇校招/鸡汤文都会出现，不能单独作为进模型的理由（回收站源文已验证）。
_WEAK_ALONE = {
    "产品",
    "市场",
    "行政",
    "营销",
    "媒体",
    "职能",
    "策划",
    "广告",
    "编辑",
    "行业报",
    "运营",
    "运营岗",
    "宣传岗",
    "管培",
}

_ADVICE_TITLE_RE = re.compile(
    r"该怎么办|人物丨|毕业典礼|转专业机制|生涯发展咨询|假期躺平|漫谈"
)

_TECH_JOB_RE = re.compile(
    r"软件开发|软件工程师|后端开发|后端工程师|前端开发|算法工程师|机器学习工程师|"
    r"人工智能工程师|人工智能(?!产品)|机器人工程师|机器人研发|"
    r"嵌入式|芯片|Java工程师|Java开发|C\+\+|iOS开发|安卓开发|Android开发|"
    r"测试开发|运维工程师|数据开发|大数据工程师|硬件研发|结构工程师|"
    r"机械设计|自动化工程师"
)


def _lexicon_path(config: Optional[dict] = None) -> str:
    custom = ""
    if config:
        custom = str(config.get("relevance_lexicon") or "").strip()
    if custom:
        return paths._resolve_under_root(custom)
    return os.path.join(paths.PROJECT_ROOT, _LEXICON_REL)


def load_lexicon(config: Optional[dict] = None) -> dict[str, Any]:
    path = _lexicon_path(config)
    if not os.path.isfile(path):
        return {"majors": [], "job_titles": [], "extra_hints": []}
    with open(path, encoding="utf-8") as f:
        data = json.load(f)
    if not isinstance(data, dict):
        return {"majors": [], "job_titles": [], "extra_hints": []}
    return data


def _norm_term(s: str) -> str:
    return re.sub(r"\s+", "", (s or "").strip())


def _iter_csv_titles(csv_path: str) -> Iterable[str]:
    if not os.path.isfile(csv_path):
        return
    with open(csv_path, encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        for row in reader:
            name = (row.get("positionName") or row.get("岗位名称") or "").strip()
            if name and _TITLE_KEEP.search(name):
                yield name


def _iter_xlsx_titles(xlsx_path: str) -> Iterable[str]:
    if not os.path.isfile(xlsx_path):
        return
    try:
        from openpyxl import load_workbook
    except ImportError:
        return
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    sheet = None
    for name in wb.sheetnames:
        if "岗位" in name and "汇总" in name:
            sheet = wb[name]
            break
    if sheet is None:
        sheet = wb[wb.sheetnames[0]]
    header = None
    col = None
    for i, row in enumerate(sheet.iter_rows(values_only=True)):
        vals = ["" if c is None else str(c).strip() for c in row]
        if i == 0:
            header = vals
            for idx, h in enumerate(header):
                if h in ("岗位名称", "岗位名", "positionName", "岗位", "具体岗位方向", "岗位方向"):
                    col = idx
                    break
            if col is None:
                col = 0
            continue
        if col is not None and col < len(vals) and vals[col]:
            yield vals[col]
    wb.close()


def collect_terms(config: Optional[dict] = None) -> list[str]:
    lex = load_lexicon(config)
    terms: list[str] = []
    for key in ("majors", "job_titles", "extra_hints"):
        for item in lex.get(key) or []:
            t = _norm_term(str(item))
            if len(t) >= _MIN_TITLE_LEN:
                terms.append(t)
    extra_csv = ""
    extra_xlsx = ""
    if config:
        extra_csv = str(config.get("relevance_titles_csv") or "").strip()
        extra_xlsx = str(config.get("relevance_titles_xlsx") or "").strip()
    default_csv = paths.csv_path() if hasattr(paths, "csv_path") else ""
    for src in (extra_csv, default_csv if extra_csv == "" else ""):
        if not src:
            continue
        path = paths._resolve_under_root(src) if not os.path.isabs(src) else src
        for name in _iter_csv_titles(path):
            t = _norm_term(name)
            if len(t) >= _MIN_TITLE_LEN:
                terms.append(t)
    if extra_xlsx:
        path = paths._resolve_under_root(extra_xlsx) if not os.path.isabs(extra_xlsx) else extra_xlsx
        for name in _iter_xlsx_titles(path):
            t = _norm_term(name)
            if _TITLE_KEEP.search(name) or any(m in name for m in (lex.get("majors") or [])):
                terms.append(_norm_term(name))
    uniq: list[str] = []
    seen: set[str] = set()
    for t in sorted(set(terms), key=lambda x: (-len(x), x)):
        if t in seen or len(t) < _MIN_TITLE_LEN:
            continue
        seen.add(t)
        uniq.append(t)
    return uniq


def match_text(text: str, terms: Optional[list[str]] = None, config: Optional[dict] = None) -> dict[str, Any]:
    blob = _norm_term(text or "")
    if terms is None:
        terms = collect_terms(config)
    hits = [t for t in terms if t and t in blob]
    min_hits = 1
    if config is not None and config.get("prefilter_min_hits") is not None:
        try:
            min_hits = max(1, int(config.get("prefilter_min_hits")))
        except (TypeError, ValueError):
            min_hits = 1
    return {
        "ok": len(hits) >= min_hits,
        "hits": hits[:20],
        "hit_count": len(hits),
        "term_count": len(terms),
    }


def is_irrelevant_tech_job(position_name: str, job_desc: object = "") -> bool:
    """岗位名是软件开发/人工智能/机器人等纯技术岗，且不像内容/产品/传媒岗。"""
    name = str(position_name or "")
    if not name.strip():
        return False
    dump = ("、" in name) or ("|" in name) or (len(name) > 40 and name.count("，") >= 2)
    if dump and re.search(r"软件开发|算法工程师|机器人|人工智能|机器学习|后端开发|前端开发", name):
        return True
    if re.search(r"产品经理|产品运营|记者|编辑|编导|融媒体|新媒体运营|管培生|宣传岗|公关", name):
        return False
    if _TECH_JOB_RE.search(name):
        return True
    if re.search(r"机器人|人工智能|算法工程师|软件开发", name):
        return True
    return False


def _meaningful_hits(hits: list[str]) -> list[str]:
    return [h for h in hits if h and h not in _WEAK_ALONE]


def should_extract(md_text: str, title: str = "", config: Optional[dict] = None) -> dict[str, Any]:
    if config is None:
        config = {}
    if config.get("prefilter_enabled", True) is False:
        return {"ok": True, "hits": [], "hit_count": 0, "skipped": False, "reason": "disabled"}
    if title and _ADVICE_TITLE_RE.search(title):
        return {
            "ok": False,
            "hits": [],
            "hit_count": 0,
            "skipped": True,
            "reason": "生涯鸡汤/人物稿，不进模型",
        }
    blob = f"{title}\n{md_text}"
    result = match_text(blob, config=config)
    useful = _meaningful_hits(result.get("hits") or [])
    strong = bool(_STRONG_RE.search(blob))
    tech = bool(_TECH_RE.search(blob))
    keep_article = bool(_KEEP_ARTICLE_RE.search(blob))
    if tech and not strong and not keep_article:
        result["ok"] = False
        result["skipped"] = True
        result["reason"] = "技术岗导向（机器人/人工智能/软件开发等）且无新闻传播强信号"
        return result
    if not useful and not strong and not keep_article:
        result["ok"] = False
        result["skipped"] = True
        result["reason"] = "仅命中过宽词（产品/市场/媒体/编辑等），不进模型"
        return result
    result["ok"] = bool(useful or strong or keep_article)
    result["skipped"] = not result["ok"]
    result["reason"] = "" if result["ok"] else "关键词预筛未命中"
    return result


def main(argv: Optional[list[str]] = None) -> int:
    parser = argparse.ArgumentParser(description="抽岗前关键词预筛（不调用 LLM）")
    parser.add_argument("--file", help="单个 Markdown")
    parser.add_argument("--dir", help="递归统计目录下 .md 命中率")
    parser.add_argument("--xlsx", help="从信息源 xlsx 岗位汇总追加岗位名（仅打印将加载的词数）")
    args = parser.parse_args(argv)
    cfg: dict[str, Any] = {}
    if args.xlsx:
        cfg["relevance_titles_xlsx"] = args.xlsx
    terms = collect_terms(cfg)
    print(f"词表 {len(terms)} 条")
    if args.file:
        with open(args.file, encoding="utf-8") as f:
            text = f.read()
        r = should_extract(text, config=cfg)
        print(json.dumps(r, ensure_ascii=False, indent=2))
        return 0 if r["ok"] else 2
    if args.dir:
        total = hit = 0
        for root, _, files in os.walk(args.dir):
            for name in files:
                if not name.lower().endswith(".md"):
                    continue
                total += 1
                path = os.path.join(root, name)
                with open(path, encoding="utf-8") as f:
                    text = f.read()
                if should_extract(text, config=cfg)["ok"]:
                    hit += 1
        miss = total - hit
        print(f"文件 {total}，预筛通过 {hit}，跳过 {miss}，预估可少调用 LLM {miss} 次")
        return 0
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
