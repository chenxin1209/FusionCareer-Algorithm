"""读入他人已结构化的岗位表，映射到 JobPost 字段名后再筛选。"""

from __future__ import annotations

import csv
import json
import os
from typing import Any, Optional

from job_structuring.engine import format_position_row
from job_structuring.enums import POSITION_FIELDNAMES, ZH_TO_CAMEL

# 朋友表头常见别名 → JobPost camelCase（只补筛选真正用到的 + 常见列）
_FIELD_ALIASES: dict[str, tuple[str, ...]] = {
    "positionName": (
        "positionName",
        "position_name",
        "job_title",
        "jobTitle",
        "position",
        "title",
        "岗位名称",
        "岗位",
        "职位名称",
        "职位",
        "岗位名",
    ),
    "companyName": (
        "companyName",
        "company_name",
        "company",
        "单位名称",
        "单位",
        "公司名称",
        "公司",
        "用人单位",
    ),
    "department": ("department", "部门", "院系"),
    "reqMajor": (
        "reqMajor",
        "req_major",
        "major",
        "专业要求",
        "专业",
        "招聘专业",
    ),
    "jobDesc": (
        "jobDesc",
        "job_desc",
        "description",
        "jd",
        "岗位描述",
        "职位描述",
        "工作内容",
        "职责",
    ),
    "reqSkills": ("reqSkills", "req_skills", "skills", "技能要求", "任职要求"),
    "reqOther": ("reqOther", "req_other", "其他要求", "投递说明"),
    "workCity": ("workCity", "work_city", "city", "工作城市", "工作地点", "城市"),
    "recruitType": ("recruitType", "recruit_type", "招聘类型"),
    "jobCategory": ("jobCategory", "job_category", "岗位大类"),
    "sourceUrl": ("sourceUrl", "source_url", "url", "link", "来源链接"),
}


def _norm_header(name: str) -> str:
    return (name or "").strip().lstrip("\ufeff")


def _lookup(row: dict, aliases: tuple[str, ...]) -> str:
    lower_map = {_norm_header(str(k)).lower(): v for k, v in row.items()}
    for alias in aliases:
        if alias in row and str(row.get(alias) or "").strip():
            return str(row.get(alias)).strip()
        hit = lower_map.get(alias.lower())
        if hit is not None and str(hit).strip():
            return str(hit).strip()
    return ""


def adapt_external_row(raw: dict) -> dict:
    """任意表头的一行 → 可交给 should_keep_job / format_position_row 的 dict。"""
    mapped = dict(raw)
    for zh, camel in ZH_TO_CAMEL.items():
        if zh in raw and camel not in mapped:
            mapped[camel] = raw[zh]
    for camel, aliases in _FIELD_ALIASES.items():
        if str(mapped.get(camel) or "").strip():
            continue
        val = _lookup(raw, aliases)
        if val:
            mapped[camel] = val
    out = format_position_row(mapped)
    oid = raw.get("id") or mapped.get("id")
    if oid is not None and str(oid).strip():
        out["id"] = str(oid).strip()
    return out


def load_structured_any(path: str) -> list[dict]:
    """支持 .csv / .json / .xlsx，返回已映射的岗位行。"""
    path = os.path.abspath(path)
    if not os.path.isfile(path):
        raise FileNotFoundError(f"未找到文件: {path}")

    ext = os.path.splitext(path)[1].lower()
    if ext == ".json":
        raw_rows = _load_json(path)
    elif ext in {".xlsx", ".xlsm"}:
        raw_rows = _load_xlsx(path)
    elif ext in {".csv", ".tsv"}:
        raw_rows = _load_csv(path, delimiter="\t" if ext == ".tsv" else ",")
    else:
        raise ValueError(f"暂不支持的格式: {ext}（请用 csv / json / xlsx）")

    return [adapt_external_row(r) for r in raw_rows if isinstance(r, dict)]


def _load_json(path: str) -> list[dict]:
    with open(path, "r", encoding="utf-8") as f:
        data = json.load(f)
    if isinstance(data, list):
        return [x for x in data if isinstance(x, dict)]
    if isinstance(data, dict):
        for key in ("jobs", "positions", "data", "records", "items"):
            v = data.get(key)
            if isinstance(v, list):
                return [x for x in v if isinstance(x, dict)]
        return [data]
    raise ValueError(f"JSON 顶层无法识别为岗位列表: {path}")


def _load_csv(path: str, delimiter: str = ",") -> list[dict]:
    with open(path, "r", encoding="utf-8-sig", newline="") as f:
        return [dict(r) for r in csv.DictReader(f, delimiter=delimiter)]


def _load_xlsx(path: str) -> list[dict]:
    try:
        from openpyxl import load_workbook
    except ImportError as e:
        raise ImportError("读取 Excel 请先: pip install openpyxl") from e
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    header_row = next(rows_iter, None)
    if not header_row:
        return []
    headers = [_norm_header(str(h) if h is not None else "") for h in header_row]
    out: list[dict] = []
    for values in rows_iter:
        if values is None or all(v is None or str(v).strip() == "" for v in values):
            continue
        item = {}
        for i, h in enumerate(headers):
            if not h:
                continue
            item[h] = values[i] if i < len(values) else ""
        out.append(item)
    return out


def write_rows_csv(path: str, rows: list[dict]) -> None:
    parent = os.path.dirname(path)
    if parent:
        os.makedirs(parent, exist_ok=True)
    fieldnames = list(POSITION_FIELDNAMES)
    extra = []
    for row in rows:
        for k in row:
            if k not in fieldnames and k not in extra and not str(k).startswith("_"):
                extra.append(k)
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fieldnames + extra, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)


def write_rows_json(path: str, rows: list[dict]) -> None:
    parent = os.path.dirname(path)
    if parent:
        os.makedirs(parent, exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        json.dump(rows, f, ensure_ascii=False, indent=2)


def detect_columns(rows: list[dict]) -> dict[str, Any]:
    """方便你确认朋友的表头有没有对上岗位名 / 专业。"""
    if not rows:
        return {"positionName": None, "reqMajor": None, "companyName": None, "keysInFile": []}
    sample = rows[0]
    return {
        "positionName": bool(sample.get("positionName")),
        "reqMajor": bool(sample.get("reqMajor")),
        "companyName": bool(sample.get("companyName")),
        "keysInFile": sorted(str(k) for k in sample.keys()),
    }
