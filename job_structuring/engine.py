"""
从 Markdown 招聘推文中用 LLM（如 DeepSeek API，模型 deepseek-chat 等）抽取结构化字段，
写入 all_positions.csv。支持「一文多岗」：模型输出 JSON 数组，每条对应一个岗位。

LLM 侧使用【中文】字段名与枚举文案（与后端 fusioncareer-api 中 desc 对齐）；落盘 CSV 仍为
camelCase + 英文枚举常量名，便于直接对接 JobPostRequest。

与 wechat_crawler 解耦：爬虫只存 md，本脚本单独运行做单篇、批量或导出。
单条岗位缺少单位名称或岗位名称则跳过该条；若整篇无任何有效岗位则写入 non_recruitment.log。

去重规则（写入 CSV 前）：
  键 = (来源链接或来源文件, 单位名称, 岗位名称)，比较前做规范化（忽略大小写与多余空白）。
  同一篇文章 LLM 返回的重复岗位、CSV 已有记录、批量运行中本次已写入记录均会跳过。
  可用 --dedup-csv 对已有 all_positions.csv 全量去重（保留首次出现行）。
"""
import csv
import json
import os
import re
import sys
from typing import Any, Optional, Set, Tuple

import requests

from job_structuring import paths

CONFIG_FILE = "config.json"
CSV_FILE = "all_positions.csv"
XLSX_FILE = "all_positions.xlsx"
SKIP_LOG = "non_recruitment.log"

# 与 JobPostRequest 对齐的 camelCase 列 + 溯源（不含 id/createdBy 等仅后端字段）
POSITION_FIELDNAMES = (
    "sourceType",
    "sourceUrl",
    "companyName",
    "department",
    "positionName",
    "jobCategory",
    "jobSubCategory",
    "recruitType",
    "headcount",
    "workStartDate",
    "workEndDate",
    "workDaysPerWeek",
    "workDurationType",
    "workPeriodType",
    "workMode",
    "workCity",
    "workProvince",
    "workLocation",
    "salaryMin",
    "salaryMax",
    "salaryDisplay",
    "jobDesc",
    "reqEduLevel",
    "reqMajor",
    "reqGradYear",
    "reqSkills",
    "reqOther",
    "status",
    "source_md",
)

XLSX_HEADER_LABELS = (
    "来源类型",
    "来源链接",
    "单位名称",
    "部门",
    "岗位名称",
    "岗位大类",
    "岗位二级分类",
    "招聘类型",
    "招聘人数",
    "工作开始日",
    "工作结束日",
    "每周工作天数",
    "每周工作天数类型",
    "实习总时长类型",
    "工作形式",
    "工作城市",
    "工作省份",
    "工作地点原文",
    "薪资下限",
    "薪资上限",
    "薪资展示",
    "岗位描述",
    "学历要求",
    "专业要求",
    "届别要求",
    "技能要求",
    "其他要求/投递说明",
    "岗位状态",
    "来源 Markdown",
)

# 中文表头字段 -> CSV / JobPost camelCase（LLM 输出中文键名）
ZH_TO_CAMEL = {
    "来源类型": "sourceType",
    "来源链接": "sourceUrl",
    "单位名称": "companyName",
    "部门": "department",
    "岗位名称": "positionName",
    "岗位大类": "jobCategory",
    "岗位二级分类": "jobSubCategory",
    "招聘类型": "recruitType",
    "招聘人数": "headcount",
    "工作开始日": "workStartDate",
    "工作结束日": "workEndDate",
    "每周工作天数": "workDaysPerWeek",
    "每周工作天数类型": "workDurationType",
    "实习总时长类型": "workPeriodType",
    "工作形式": "workMode",
    "工作城市": "workCity",
    "工作省份": "workProvince",
    "工作地点原文": "workLocation",
    "薪资下限": "salaryMin",
    "薪资上限": "salaryMax",
    "薪资展示": "salaryDisplay",
    "岗位描述": "jobDesc",
    "学历要求": "reqEduLevel",
    "专业要求": "reqMajor",
    "届别要求": "reqGradYear",
    "技能要求": "reqSkills",
    "其他要求与投递说明": "reqOther",
    "岗位状态": "status",
}

# 后端枚举中文 desc -> Java 枚举常量（与 fusioncareer-api 一致）
_ZH_JOB_CATEGORY = {
    "学术教职": "ACADEMIC",
    "党政机关": "GOVERNMENT",
    "新闻媒体": "MEDIA",
    "企业公司": "ENTERPRISE",
    "企业": "ENTERPRISE",
    "其他": "OTHER",
    " 其他": "OTHER",
    # 易错：国际组织/研究院属「党政机关-其他事业单位」，不是「其他」
    "国际组织": "GOVERNMENT",
    "事业单位": "GOVERNMENT",
}
_ZH_JOB_SUB_CATEGORY = {
    "升学深造": "FURTHER_STUDY",
    "考取教职": "TEACHING_POSITION",
    "中学教师": "MIDDLE_SCHOOL_TEACHER",
    "选调生": "SELECTED_GRADUATE",
    "公务员": "CIVIL_SERVANT",
    "高校行政": "UNIVERSITY_ADMIN",
    "医院": "HOSPITAL",
    "银行": "BANK",
    "其他事业单位": "OTHER_PUBLIC_INSTITUTION",
    "党报央媒": "CENTRAL_MEDIA",
    "地区主流媒体": "REGIONAL_MEDIA",
    "其他媒体机构": "OTHER_MEDIA",
    "自媒体": "SELF_MEDIA",
    "国央企": "STATE_OWNED",
    "民企": "PRIVATE_ENTERPRISE",
    "外企": "FOREIGN_ENTERPRISE",
    "其他": "OTHER",
    # 常见 LLM 非标准表述 -> 其他事业单位 / 企业
    "国际组织": "OTHER_PUBLIC_INSTITUTION",
    "国际机构": "OTHER_PUBLIC_INSTITUTION",
    "联合国": "OTHER_PUBLIC_INSTITUTION",
    "联合国系统": "OTHER_PUBLIC_INSTITUTION",
    "研究院": "OTHER_PUBLIC_INSTITUTION",
    "科研院所": "OTHER_PUBLIC_INSTITUTION",
    "实验室": "OTHER_PUBLIC_INSTITUTION",
    "民营企业": "PRIVATE_ENTERPRISE",
    "私营企业": "PRIVATE_ENTERPRISE",
    "外资企业": "FOREIGN_ENTERPRISE",
    "国有企业": "STATE_OWNED",
    "央企": "STATE_OWNED",
}
_ZH_RECRUIT_TYPE = {
    "大实习": "BIG_INTERNSHIP",
    "小实习": "SMALL_INTERNSHIP",
    "日常实习": "DAILY_INTERNSHIP",
    "应届生招聘": "CAMPUS_RECRUITMENT",
    "应届生摸排": "CAMPUS_SCREENING",
    "其他": "OTHER",
}
_ZH_WORK_DURATION = {
    "一周1-2天": "ONE_TO_TWO_DAYS",
    "一周3-4天": "THREE_TO_FOUR_DAYS",
    "一周5天": "FIVE_DAYS",
}
_ZH_WORK_PERIOD = {
    "3个月以内": "LESS_THAN_THREE_MONTHS",
    "3-6个月": "THREE_TO_SIX_MONTHS",
    "6个月以上": "MORE_THAN_SIX_MONTHS",
}
_ZH_WORK_MODE = {
    "线上": "ONLINE",
    "线下": "OFFLINE",
    "线上线下均可": "HYBRID",
}
_ZH_EDU_LEVEL = {
    "本科生": "UNDERGRADUATE",
    "学术硕士研究生": "ACADEMIC_MASTER",
    "专业硕士研究生": "PROFESSIONAL_MASTER",
    "硕士研究生": "ACADEMIC_MASTER",
    "硕士": "ACADEMIC_MASTER",
    "博士研究生": "DOCTORAL",
    "博士": "DOCTORAL",
}
_ZH_JOB_STATUS = {
    "已下线": "OFFLINE",
    "发布中": "PUBLISHED",
    "已截止": "EXPIRED",
}
_ZH_SOURCE_TYPE = {
    "平台发布": "PLATFORM",
    "就业资讯源爬取": "CRAWL",
}

# 允许模型直接输出英文枚举常量时的合法集合（小校验）
_VALID = {
    "sourceType": set(_ZH_SOURCE_TYPE.values()) | {"PLATFORM", "CRAWL"},
    "jobCategory": set(_ZH_JOB_CATEGORY.values()),
    "jobSubCategory": set(_ZH_JOB_SUB_CATEGORY.values()) | {"OTHER"},
    "recruitType": set(_ZH_RECRUIT_TYPE.values()),
    "workDurationType": set(_ZH_WORK_DURATION.values()),
    "workPeriodType": set(_ZH_WORK_PERIOD.values()),
    "workMode": set(_ZH_WORK_MODE.values()),
    "reqEduLevel": set(_ZH_EDU_LEVEL.values()),
    "status": set(_ZH_JOB_STATUS.values()),
}

# JobSubCategory 常量 -> 应对的 JobCategory（用于修正大类/子类不一致）
JOB_SUB_TO_PARENT = {
    "FURTHER_STUDY": "ACADEMIC",
    "TEACHING_POSITION": "ACADEMIC",
    "MIDDLE_SCHOOL_TEACHER": "ACADEMIC",
    "SELECTED_GRADUATE": "GOVERNMENT",
    "CIVIL_SERVANT": "GOVERNMENT",
    "UNIVERSITY_ADMIN": "GOVERNMENT",
    "HOSPITAL": "GOVERNMENT",
    "BANK": "GOVERNMENT",
    "OTHER_PUBLIC_INSTITUTION": "GOVERNMENT",
    "CENTRAL_MEDIA": "MEDIA",
    "REGIONAL_MEDIA": "MEDIA",
    "OTHER_MEDIA": "MEDIA",
    "SELF_MEDIA": "MEDIA",
    "STATE_OWNED": "ENTERPRISE",
    "PRIVATE_ENTERPRISE": "ENTERPRISE",
    "FOREIGN_ENTERPRISE": "ENTERPRISE",
    "OTHER": "OTHER",
}

LLM_SYSTEM_PROMPT = """Role: 你是校园招聘结构化抽取助手。
Task: 阅读 Markdown 推文，识别其中【每一个】独立的招聘岗位（一文多岗须拆成多条），输出 **一个 JSON 数组**（仅数组，不要外层对象）。

## 输出格式（硬性）
1. 顶层必须是 JSON List：`[ {...}, {...} ]`。
2. 数组中每个元素是一个岗位对象；字段名必须使用下列【中文】键名（不要 camelCase）。
3. 若文章开头/导语存在适用于全文的「报名截止日期」「统一投递方式/邮箱/网申链接」等，请把该信息写入【每一个】岗位对象的对应字段；若某一条岗位段落里有更具体的日期或投递方式，以该条为准覆盖通用值。
4. 枚举类字段请填【中文】，且必须与下方「枚举可选中文」完全一致（不要用英文常量名，不要填 code 数字）。
5. 数字字段：招聘人数、每周工作天数、薪资下限、薪资上限用 JSON 数字，不详用 null。
6. 日期字段用字符串 `YYYY-MM-DD`，不详用空字符串 ""。
7. 非招聘、无法拆出任何有效岗位时输出空数组 `[]`。

## 每个岗位对象须包含的键（可空串或 null，但键名要齐全）
- 来源类型：固定填「就业资讯源爬取」
- 来源链接：可空（脚本会从 Markdown 注入原文链接）
- 单位名称
- 部门
- 岗位名称
- 岗位大类
- 岗位二级分类
- 招聘类型
- 招聘人数
- 工作开始日
- 工作结束日
- 每周工作天数
- 每周工作天数类型
- 实习总时长类型
- 工作形式
- 工作城市
- 工作省份
- 工作地点原文
- 薪资下限
- 薪资上限
- 薪资展示
- 岗位描述
- 学历要求
- 专业要求
- 届别要求
- 技能要求
- 其他要求与投递说明
- 岗位状态：新抽取且仍有效填「发布中」；正文明确已截止填「已截止」

## 枚举可选中文（须严格一致）
- 来源类型：平台发布 | 就业资讯源爬取
- 岗位大类：学术教职 | 党政机关 | 新闻媒体 | 企业公司 | 其他
- 岗位二级分类：升学深造 | 考取教职 | 中学教师 | 选调生 | 公务员 | 高校行政 | 医院 | 银行 | 其他事业单位 | 党报央媒 | 地区主流媒体 | 其他媒体机构 | 自媒体 | 国央企 | 民企 | 外企 | 其他
- 招聘类型：大实习 | 小实习 | 日常实习 | 应届生招聘 | 应届生摸排 | 其他
- 每周工作天数类型：一周1-2天 | 一周3-4天 | 一周5天
- 实习总时长类型：3个月以内 | 3-6个月 | 6个月以上
- 工作形式：线上 | 线下 | 线上线下均可
- 学历要求：本科生 | 学术硕士研究生 | 专业硕士研究生 | 硕士研究生 | 博士研究生（不限可填 ""）
- 岗位状态：已下线 | 发布中 | 已截止

## 岗位大类与二级分类（判定指南，优先遵守，慎填「其他」）
**学术教职** + 二级：高校教职岗、博后、中学教师、升学深造（留学/保研项目）等。
**党政机关** + 二级：
  - 选调生、公务员、高校行政（辅导员/行政岗）；
  - 医院、银行；
  - **其他事业单位**：科研院所/研究院/实验室/勘测设计院、协会学会、智库、**联合国及国际组织**（WMO、WHO、世行、驻华使领馆系统、国际民航组织等）、出版社（事业编制）等——凡政府机关以外的事业性质单位均归此类，**不要**把国际组织标成「其他」大类。
**新闻媒体** + 二级：党报央媒、地区主流媒体、其他媒体机构、自媒体/MCN。
**企业公司** + 二级：
  - **国央企**：中央/地方国资委体系、央企、国有大行/政策性银行总行、中国XX集团/研究院（企业法人）等；
  - **民企**：民营有限公司/股份公司、科技创业公司、律所（合伙制企业）等；
  - **外企**：外商独资、合资企业中明确外资品牌（如宝洁、玛氏、四大会计外资所等）。
  凡招聘主体为「XX有限公司/股份公司/集团」且非明显央企的，默认**民企**，不要填「其他」。
**「其他」大类**仅用于：无法归入以上四类的极少数情况；**禁止**因懒得判断而填「其他」。

## 不应抽取为岗位记录的内容（须输出空数组 `[]`）
- 校园**双选会/招聘会/专场活动**通知、邀请函、展位图、参会指南：仅有时间地点、无具体用人单位与岗位名称的；
- 多校联合招聘会通稿、宣讲会周历/预告（无单场企业岗位详情）；
- 就业政策宣传、生涯咨询预约、讲座培训、赛事获奖、校园文化活动、毕业手续指南等与**具体用人单位招聘**无关的推文。

## 可抽取的情形
- 文章含**具体单位 + 具体岗位/实习岗位名称**（含联合国实习、研究院招聘、民企校招等），按上文分类；
- 招聘会通稿文末附**带企业名与岗位名**的招聘清单：按企业拆条，分别归类二级分类。

Rules:
- 仅输出 JSON 数组本身，不要 markdown 代码围栏外的解释文字。
- 「多企业汇总」类文章，每个参展/招聘单位拆成独立岗位对象（单位名称不同）。
- 岗位大类、岗位二级分类必须成对匹配（如：企业公司+民企，党政机关+其他事业单位）；勿只填大类不填二级。
"""


# 旧版 LLM / CSV 字段 -> JobPost 字段
_LEGACY_LLM_TO_CAMEL = (
    ("companyName", ("company_name", "company")),
    ("positionName", ("job_title", "position")),
    ("workCity", ("work_city", "city")),
    ("workEndDate", ("deadline",)),
    ("reqSkills", ("requirement_summary",)),
    ("reqOther", ("application_method",)),
)


def _config_path():
    return paths.config_path()


def load_json(filepath):
    if os.path.exists(filepath):
        with open(filepath, "r", encoding="utf-8") as f:
            try:
                return json.load(f)
            except json.JSONDecodeError:
                return {}
    return {}


def load_config():
    return load_json(_config_path())


def _article_link_from_md(text: str) -> str:
    m = re.search(r"^\*\*Link:\*\*\s*(.+)$", text, re.MULTILINE)
    return m.group(1).strip() if m else ""


def _article_title_from_md(text: str) -> str:
    m = re.search(r"^#\s+(.+)$", text, re.MULTILINE)
    return m.group(1).strip() if m else ""


def _article_title_from_path(md_path: str) -> str:
    base = os.path.basename(md_path)
    m = re.match(r"\d{4}-\d{2}-\d{2}[_\s](.+)\.md$", base, re.I)
    return m.group(1).strip() if m else os.path.splitext(base)[0]


def _resolve_article_title(md_text: str, md_path: str) -> str:
    return _article_title_from_md(md_text) or _article_title_from_path(md_path)


# 标题过滤：不含具体招聘信号时命中则跳过 LLM
_NON_RECRUITMENT_TITLE_HINTS = (
    "就业手续",
    "毕业去向",
    "就业咨询",
    "生涯咨询",
    "职业咨询",
    "咨询预约",
    "咨询工作室",
    "春招小课堂",
    "创业云",
    "政策每周汇",
    "政策清单",
    "支持政策",
    "精彩回顾",
    "活动回顾",
    "国赛",
    "获奖",
    "表彰",
    "倡议书",
    "五四",
    "网络面试间",
    "就业加油站",
    "职行力工作室",
    "受聘仪式",
    "人才交流会举行",
    "赛事首金",
    "文化月",
    "假期躺平",
    "漫谈",
)
_JOB_FAIR_TITLE_HINTS = (
    "双选会",
    "双选",
    "招聘会",
    "专场招聘",
    "宣讲会预告",
    "宣讲会安排",
    "展位图",
    "用人单位邀请函",
    "学生邀请函",
    "参会指南",
    "线上招聘会",
    "春季招聘会",
    "秋季招聘会",
    "招聘会场",
    "招聘会｜",
    "招聘会|",
)


def _title_implies_concrete_hiring(title: str) -> bool:
    """标题是否表明文章含具体用人单位招聘/实习信息。"""
    if not title:
        return False
    if re.search(r"【\s*(招聘|实习|工作|岗位|事业|国际组织|央企|国企|26届|27届)", title):
        return True
    if re.search(
        r"(校园招聘|人才招聘|公开招聘|招聘公告|招聘启事|实习招聘|管培生|【\d{2}届)",
        title,
    ):
        return True
    if re.search(r"(^|[\|｜\s])招聘[\s｜|]|^实习[\s｜|]|就业实习\s+招聘|招聘信息\d", title):
        return True
    return False


def _title_skip_reason(title: str) -> Optional[str]:
    """若应跳过本篇，返回原因文案；否则 None。"""
    if not title:
        return None
    if any(k in title for k in ("宣讲会预告", "宣讲会安排", "宣讲会周历", "宣讲会早知道")):
        return "宣讲会预告"
    if _title_implies_concrete_hiring(title):
        return None
    for kw in _NON_RECRUITMENT_TITLE_HINTS:
        if kw in title:
            return f"非就业招聘({kw})"
    for kw in _JOB_FAIR_TITLE_HINTS:
        if kw in title:
            return f"招聘会/双选会({kw})"
    if re.search(r"活动(预告|报名)", title) and "招聘" not in title:
        return "非招聘活动"
    return None


def _log_skip_article(md_path: str, reason: str, title: str = "") -> None:
    label = f"{reason}: {title} | {md_path}" if title else f"{reason}: {md_path}"
    print(f"  [structure_data] 非招聘信息(标题过滤-{label})")
    log_path = paths.skip_log_path()
    os.makedirs(os.path.dirname(log_path) or ".", exist_ok=True)
    with open(log_path, "a", encoding="utf-8") as logf:
        logf.write(f"非招聘信息(标题过滤-{label})\n")


def _extract_json_object(raw: str) -> Optional[dict]:
    """解析单个 JSON 对象（兼容旧版 LLM）。"""
    raw = raw.strip()
    m = re.search(r"```(?:json)?\s*([\s\S]*?)```", raw, re.IGNORECASE)
    if m:
        raw = m.group(1).strip()
    try:
        obj = json.loads(raw)
        return obj if isinstance(obj, dict) else None
    except json.JSONDecodeError:
        pass
    m2 = re.search(r"\{[\s\S]*\}", raw)
    if m2:
        try:
            obj = json.loads(m2.group(0))
            return obj if isinstance(obj, dict) else None
        except json.JSONDecodeError:
            return None
    return None


def _merge_zh_common_into_position(pos: dict, common: dict) -> dict:
    """将「文章通用信息」中的键合并到岗位对象（仅填补岗位侧为空的键）。"""
    out = dict(pos)
    if not isinstance(common, dict):
        return out
    for k, v in common.items():
        if v is None:
            continue
        if isinstance(v, str) and not v.strip():
            continue
        cur = out.get(k)
        empty = cur is None or (isinstance(cur, str) and not str(cur).strip())
        if empty:
            out[k] = v
    return out


def _extract_positions_list(raw: str) -> Optional[list]:
    """
    解析 LLM 返回：优先 JSON 数组；兼容带「岗位列表」的单对象包裹；兼容旧版单对象。
    返回 dict 列表（每项为中文键或混排键的岗位对象）。
    """
    raw0 = raw.strip()
    m = re.search(r"```(?:json)?\s*([\s\S]*?)```", raw0, re.IGNORECASE)
    payload = m.group(1).strip() if m else raw0
    data = None
    try:
        data = json.loads(payload)
    except json.JSONDecodeError:
        m2 = re.search(r"\[[\s\S]*\]", payload)
        if m2:
            try:
                data = json.loads(m2.group(0))
            except json.JSONDecodeError:
                data = None
        if data is None:
            m3 = re.search(r"\{[\s\S]*\}", payload)
            if m3:
                try:
                    data = json.loads(m3.group(0))
                except json.JSONDecodeError:
                    return None
    if isinstance(data, list):
        return [dict(x) for x in data if isinstance(x, dict)]
    if isinstance(data, dict):
        for k in ("岗位列表", "positions", "jobs"):
            v = data.get(k)
            if isinstance(v, list):
                common = data.get("文章通用信息") or data.get("通用信息") or {}
                out_list = []
                for item in v:
                    if not isinstance(item, dict):
                        continue
                    merged = _merge_zh_common_into_position(dict(item), common) if common else dict(item)
                    out_list.append(merged)
                return out_list
        return [_apply_legacy_llm_aliases(dict(data))]
    return None


def _zh_lookup(zh_map: dict, s: str) -> Optional[str]:
    s = (s or "").strip()
    if not s:
        return None
    if s in zh_map:
        return zh_map[s]
    compact = re.sub(r"\s+", "", s)
    for k, v in zh_map.items():
        if re.sub(r"\s+", "", (k or "").strip()) == compact:
            return v
    return None


def _map_enum_value(field: str, raw: object) -> str:
    """中文枚举或已是英文常量 -> 后端枚举英文名；空值按字段给默认。"""
    s = _normalize_field(raw)
    if not s:
        if field in ("jobCategory", "jobSubCategory", "recruitType"):
            return "OTHER"
        return ""
    if field in _VALID and s in _VALID[field]:
        return s
    zm = {
        "sourceType": _ZH_SOURCE_TYPE,
        "jobCategory": _ZH_JOB_CATEGORY,
        "jobSubCategory": _ZH_JOB_SUB_CATEGORY,
        "recruitType": _ZH_RECRUIT_TYPE,
        "workDurationType": _ZH_WORK_DURATION,
        "workPeriodType": _ZH_WORK_PERIOD,
        "workMode": _ZH_WORK_MODE,
        "reqEduLevel": _ZH_EDU_LEVEL,
        "status": _ZH_JOB_STATUS,
    }.get(field)
    if zm:
        hit = _zh_lookup(zm, s)
        if hit:
            return hit
    if field in ("jobCategory", "jobSubCategory", "recruitType"):
        return "OTHER"
    return ""


def _flatten_zh_record_to_camel(zh_obj: dict) -> dict:
    """中文键 -> camelCase 中间表示（值仍为原文或数字）。"""
    flat: dict = {}
    for zh_k, camel in ZH_TO_CAMEL.items():
        if zh_k not in zh_obj:
            continue
        v = zh_obj[zh_k]
        if v is None:
            continue
        if isinstance(v, str) and not v.strip():
            continue
        flat[camel] = v
    for camel in POSITION_FIELDNAMES:
        if camel in ("source_md",):
            continue
        if camel in zh_obj and zh_obj[camel] is not None:
            v = zh_obj[camel]
            if isinstance(v, str) and not v.strip():
                continue
            if camel not in flat or not str(flat.get(camel, "")).strip():
                flat[camel] = v
    return flat


def _normalize_enums_in_row(flat: dict) -> dict:
    out = dict(flat)
    for field in (
        "sourceType",
        "jobCategory",
        "jobSubCategory",
        "recruitType",
        "workDurationType",
        "workPeriodType",
        "workMode",
        "reqEduLevel",
        "status",
    ):
        if field not in out:
            continue
        out[field] = _map_enum_value(field, out.get(field))
    return out


def _infer_category_heuristics(flat: dict, article_title: str = "") -> dict:
    """
    当 LLM 未填或填「其他」时，根据单位名称/岗位/标题做规则补全。
    仅覆盖仍为 OTHER 或空的字段。
    """
    out = dict(flat)
    cat = str(out.get("jobCategory") or "").strip()
    sub = str(out.get("jobSubCategory") or "").strip()
    company = str(out.get("companyName") or "").strip()
    position = str(out.get("positionName") or "").strip()
    blob = f"{company} {position} {article_title}"

    cat_is_other = not cat or cat == "OTHER"
    sub_is_other = not sub or sub == "OTHER"
    if not cat_is_other and not sub_is_other:
        return out

    intl_markers = (
        "联合国",
        "UN ",
        "UN-",
        "WMO",
        "WHO",
        "世界银行",
        "国际组织",
        "国际民航",
        "教科文组织",
        "人居署",
        "开发计划署",
        "高专办",
        "Office of the",
        "Commission for",
        "Coordinator System",
    )
    research_markers = (
        "研究院",
        "研究所",
        "科学院",
        "实验室",
        "勘测设计",
        "融媒体中心",
    )
    state_markers = (
        "央企",
        "国有",
        "中国烟草",
        "中石油",
        "中石化",
        "国家电网",
        "中国建筑",
        "中铁",
        "中车",
        "中航",
        "招商局",
        "中粮",
        "宝武",
        "中国邮政",
    )
    media_markers = (
        "日报",
        "晚报",
        "电视台",
        "广播电视台",
        "新华社",
        "人民日报",
        "央视",
        "传媒集团",
    )

    def _apply(cat_en: str, sub_en: str) -> None:
        nonlocal cat_is_other, sub_is_other
        if cat_is_other:
            out["jobCategory"] = cat_en
        if sub_is_other:
            out["jobSubCategory"] = sub_en

    if "国际组织" in article_title or any(m in blob for m in intl_markers):
        _apply("GOVERNMENT", "OTHER_PUBLIC_INSTITUTION")
        return out
    if any(m in company for m in research_markers):
        _apply("GOVERNMENT", "OTHER_PUBLIC_INSTITUTION")
        return out
    if any(m in company for m in state_markers) or re.search(
        r"^中国[\u4e00-\u9fff]{1,10}(集团|银行|保险|证券|石油|建筑)",
        company,
    ):
        _apply("ENTERPRISE", "STATE_OWNED")
        return out
    if any(m in company for m in media_markers):
        sub_m = (
            "CENTRAL_MEDIA"
            if any(x in company for x in ("人民", "新华", "央视", "中央"))
            else "REGIONAL_MEDIA"
        )
        _apply("MEDIA", sub_m)
        return out
    if re.search(r"(有限公司|股份公司|律师事务所|科技有限)", company):
        _apply("ENTERPRISE", "PRIVATE_ENTERPRISE")
        return out

    return out


def _repair_job_category_pair(out: dict) -> dict:
    sub = out.get("jobSubCategory", "")
    cat = out.get("jobCategory", "")
    parent = JOB_SUB_TO_PARENT.get(sub)
    if parent and cat != parent:
        out["jobCategory"] = parent
    if sub == "OTHER" and not cat:
        out["jobCategory"] = "OTHER"
    if cat == "OTHER" and not sub:
        out["jobSubCategory"] = "OTHER"
    return out


def _ensure_enum_defaults(flat: dict) -> dict:
    out = dict(flat)
    if not str(out.get("jobCategory", "")).strip():
        out["jobCategory"] = "OTHER"
    if not str(out.get("jobSubCategory", "")).strip():
        out["jobSubCategory"] = "OTHER"
    if not str(out.get("recruitType", "")).strip():
        out["recruitType"] = "OTHER"
    return out


def _zh_item_to_job_row(
    zh_obj: dict,
    article_link: str,
    md_path: str,
    article_title: str = "",
) -> dict:
    flat = _flatten_zh_record_to_camel(zh_obj)
    flat = _apply_legacy_llm_aliases(flat)
    flat = _normalize_enums_in_row(flat)
    flat = _infer_category_heuristics(flat, article_title)
    flat = _repair_job_category_pair(flat)
    flat = _ensure_enum_defaults(flat)
    row: dict = {}
    for key in POSITION_FIELDNAMES:
        if key == "source_md":
            row[key] = md_path
        elif key == "sourceUrl":
            row[key] = article_link or _normalize_field(flat.get("sourceUrl"))
        elif key == "sourceType":
            row[key] = _normalize_field(flat.get("sourceType")) or "CRAWL"
        elif key == "status":
            row[key] = _normalize_field(flat.get("status")) or "PUBLISHED"
        else:
            row[key] = _normalize_field(flat.get(key))
    return row


def _norm_sig(s: str) -> str:
    return re.sub(r"\s+", " ", (s or "").strip()).lower()


def _position_dedup_key(
    source_url: str = "",
    company: str = "",
    position: str = "",
    source_md: str = "",
) -> Tuple[str, str, str]:
    """
    去重键：(来源标识, 单位, 岗位)。
    无链接时用规范化后的 source_md 路径作为来源标识，避免不同文章误合并。
    """
    u = (source_url or "").strip().lower()
    if not u and source_md:
        u = "md:" + os.path.normcase(os.path.abspath(source_md))
    co = _norm_sig(company)
    po = _norm_sig(position)
    return (u, co, po)


def _dedup_key_from_row(row: dict) -> Tuple[str, str, str]:
    return _position_dedup_key(
        row.get("sourceUrl") or row.get("article_link") or "",
        row.get("companyName", ""),
        row.get("positionName", ""),
        row.get("source_md", ""),
    )


class PositionDedupIndex:
    """内存去重索引：启动时加载 CSV，写入成功后登记新键。"""

    def __init__(self) -> None:
        self._keys: Set[Tuple[str, str, str]] = set()

    def __len__(self) -> int:
        return len(self._keys)

    def load_csv(self, csv_path: str) -> int:
        """从已有 CSV 加载键，返回加载条数。"""
        if not os.path.isfile(csv_path):
            return 0
        n = 0
        with open(csv_path, "r", encoding="utf-8-sig", newline="") as f:
            reader = csv.DictReader(f)
            for raw in reader:
                row = format_position_row(raw)
                co = row.get("companyName", "")
                po = row.get("positionName", "")
                if not co or not po:
                    continue
                key = _dedup_key_from_row(row)
                if key not in self._keys:
                    self._keys.add(key)
                    n += 1
        return n

    def contains(
        self,
        source_url: str = "",
        company: str = "",
        position: str = "",
        source_md: str = "",
    ) -> bool:
        if not _norm_sig(company) or not _norm_sig(position):
            return False
        return _position_dedup_key(source_url, company, position, source_md) in self._keys

    def contains_row(self, row: dict) -> bool:
        return _dedup_key_from_row(row) in self._keys

    def add(
        self,
        source_url: str = "",
        company: str = "",
        position: str = "",
        source_md: str = "",
    ) -> None:
        if not _norm_sig(company) or not _norm_sig(position):
            return
        self._keys.add(_position_dedup_key(source_url, company, position, source_md))

    def add_row(self, row: dict) -> None:
        self._keys.add(_dedup_key_from_row(row))


def _deduplicate_llm_items(
    items: list,
    article_link: str,
    md_path: str,
) -> Tuple[list, int]:
    """去掉同一篇 Markdown 内 LLM 返回的重复岗位（保留首次）。"""
    seen: Set[Tuple[str, str, str]] = set()
    out = []
    skipped = 0
    for item in items:
        preview = _flatten_zh_record_to_camel(item)
        company = _normalize_field(preview.get("companyName"))
        position = _normalize_field(preview.get("positionName"))
        if not company or not position:
            out.append(item)
            continue
        key = _position_dedup_key(article_link, company, position, md_path)
        if key in seen:
            skipped += 1
            continue
        seen.add(key)
        out.append(item)
    return out, skipped


def deduplicate_csv_file(csv_path: Optional[str] = None) -> Tuple[int, int]:
    """
    对 CSV 全量去重（保留首次出现的行），原地覆盖写回。
    返回 (去重前行数, 去重后行数)。
    """
    csv_path = csv_path or paths.csv_path()
    if not os.path.isfile(csv_path):
        raise FileNotFoundError(f"未找到 CSV: {csv_path}")

    rows_in = []
    with open(csv_path, "r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        fieldnames = reader.fieldnames
        for raw in reader:
            rows_in.append(format_position_row(raw))

    if fieldnames and list(fieldnames) != list(POSITION_FIELDNAMES):
        print(
            "[structure_data] 警告: CSV 表头与当前字段不一致，去重仍将按 POSITION_FIELDNAMES 写回。"
        )

    seen: Set[Tuple[str, str, str]] = set()
    rows_out = []
    dup_count = 0
    for row in rows_in:
        co = row.get("companyName", "")
        po = row.get("positionName", "")
        if not co or not po:
            rows_out.append(row)
            continue
        key = _dedup_key_from_row(row)
        if key in seen:
            dup_count += 1
            continue
        seen.add(key)
        rows_out.append(row)

    with open(csv_path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(POSITION_FIELDNAMES), extrasaction="ignore")
        w.writeheader()
        w.writerows(rows_out)

    before = len(rows_in)
    after = len(rows_out)
    print(
        f"[structure_data] CSV 去重完成: {csv_path}\n"
        f"  去重前 {before} 行，去重后 {after} 行，移除重复 {before - after} 行"
    )
    return before, after


def _llm_chat(messages: list, config: dict) -> Optional[str]:
    key = (
        config.get("llm_api_key")
        or os.environ.get("DEEPSEEK_API_KEY")
        or os.environ.get("LLM_API_KEY")
        or ""
    ).strip()
    if not key:
        print("  [structure_data] 未配置 llm_api_key，跳过 LLM（可设置环境变量 DEEPSEEK_API_KEY）")
        return None

    base = (config.get("llm_base_url") or "https://api.deepseek.com/v1").rstrip("/")
    model = config.get("llm_model") or "deepseek-chat"
    url = f"{base}/chat/completions"

    payload = {
        "model": model,
        "messages": messages,
        "temperature": 0.2,
    }
    headers = {"Authorization": f"Bearer {key}", "Content-Type": "application/json"}
    try:
        r = requests.post(url, headers=headers, json=payload, timeout=180)
        r.raise_for_status()
        data = r.json()
        return data["choices"][0]["message"]["content"]
    except Exception as e:
        print(f"  [structure_data] LLM 请求失败: {e}")
        return None


def _normalize_field(v) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, (int, float)):
        if isinstance(v, float) and v == int(v):
            return str(int(v))
        return str(v)
    s = str(v).strip()
    if s.lower() in ("null", "none", "undefined"):
        return ""
    return s


def _apply_legacy_llm_aliases(obj: dict) -> dict:
    """兼容旧版 LLM 输出的 snake_case / 旧列名。"""
    out = dict(obj)
    for canonical, legacy_keys in _LEGACY_LLM_TO_CAMEL:
        cur = str(out.get(canonical, "")).strip()
        if cur:
            continue
        for lk in legacy_keys:
            val = out.get(lk)
            if val is None or not str(val).strip():
                continue
            if canonical == "reqOther" and lk == "application_method":
                out[canonical] = f"投递方式：{_normalize_field(val)}"
            elif canonical == "reqSkills" and lk == "requirement_summary":
                out[canonical] = _normalize_field(val)
            elif canonical == "workEndDate" and lk == "deadline":
                d = _normalize_field(val)
                if d and d != "不详":
                    out[canonical] = d
            else:
                out[canonical] = _normalize_field(val)
            break

    legacy_notes = []
    for key, label in (
        ("company_type", "旧 company_type"),
        ("is_internship", "旧 is_internship"),
        ("job_category", "旧 job_category"),
    ):
        v = str(out.get(key, "")).strip()
        if v:
            legacy_notes.append(f"{label}: {v}")
    if legacy_notes:
        note = "；".join(legacy_notes)
        ro = str(out.get("reqOther", "")).strip()
        out["reqOther"] = f"{note}；{ro}" if ro else note
    return out


def merge_legacy_csv_row(row: dict) -> dict:
    """
    将旧版 CSV / 爬虫列名映射到当前 POSITION_FIELDNAMES（JobPost 对齐）。
    """
    out = dict(row)
    pairs = [
        ("companyName", ("companyName", "company_name", "company")),
        ("positionName", ("positionName", "job_title", "position")),
        ("workCity", ("workCity", "work_city", "city")),
        ("sourceUrl", ("sourceUrl", "article_link")),
        ("workEndDate", ("workEndDate", "deadline")),
        ("reqOther", ("reqOther", "application_method", "apply_link")),
        ("source_md", ("source_md",)),
    ]
    for canonical, aliases in pairs:
        if str(out.get(canonical, "")).strip():
            continue
        for a in aliases:
            if a in out and str(out.get(a, "")).strip():
                val = out.get(a)
                if canonical == "reqOther" and a == "application_method":
                    out[canonical] = f"投递方式：{_normalize_field(val)}"
                elif canonical == "reqOther" and a == "apply_link":
                    out[canonical] = f"投递链接：{_normalize_field(val)}"
                else:
                    out[canonical] = val if isinstance(val, str) else _normalize_field(val)
                break
    return out


def format_position_row(row: dict) -> dict:
    """
    统一岗位行格式：去首尾空白；普通文本折叠多余空白；
    sourceUrl、source_md 仅 strip，避免破坏链接。
    """
    row = merge_legacy_csv_row(row)
    no_collapse = {"sourceUrl", "source_md"}
    out = {}
    for key in POSITION_FIELDNAMES:
        s = _normalize_field(row.get(key))
        if key in no_collapse:
            out[key] = s
        else:
            out[key] = re.sub(r"\s+", " ", s).strip()
    if not out.get("sourceType"):
        out["sourceType"] = "CRAWL"
    if not out.get("status"):
        out["status"] = "PUBLISHED"
    if not out.get("recruitType"):
        out["recruitType"] = "OTHER"
    return out


def _append_csv_row(row: dict, csv_path: str) -> bool:
    row = format_position_row(row)
    file_exists = os.path.exists(csv_path)
    expected = list(POSITION_FIELDNAMES)
    if file_exists:
        with open(csv_path, "r", encoding="utf-8-sig", newline="") as f:
            reader = csv.DictReader(f)
            old_fields = reader.fieldnames or []
        if old_fields and list(old_fields) != expected:
            print(
                f"  [structure_data] 已有 CSV 表头与当前 JobPost 对齐列不一致，已跳过写入（请先备份后删除或重命名该 CSV，再重新抽取）。\n"
                f"    文件: {csv_path}"
            )
            return False
    with open(csv_path, "a", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=expected, extrasaction="ignore")
        if not file_exists:
            w.writeheader()
        w.writerow(row)
    return True


def export_positions_xlsx(
    csv_path: Optional[str] = None,
    xlsx_path: Optional[str] = None,
) -> str:
    """
    将 all_positions.csv 读入并导出为格式化的 xlsx（冻结首行、列宽、长文本换行）。
    返回生成的 xlsx 绝对路径。
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as e:
        raise ImportError("请先安装: pip install openpyxl") from e

    csv_path = csv_path or paths.csv_path()
    xlsx_path = xlsx_path or paths.xlsx_path()

    if not os.path.isfile(csv_path):
        raise FileNotFoundError(f"未找到 CSV: {csv_path}")

    rows = []
    with open(csv_path, "r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        for raw in reader:
            rows.append(format_position_row(raw))

    wb = Workbook()
    ws = wb.active
    ws.title = "岗位数据"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell_align = Alignment(vertical="top", wrap_text=True)

    for col, label in enumerate(XLSX_HEADER_LABELS, start=1):
        c = ws.cell(row=1, column=col, value=label)
        c.font = header_font
        c.fill = header_fill
        c.alignment = header_align

    ncols = len(POSITION_FIELDNAMES)
    for r, row in enumerate(rows, start=2):
        for col in range(1, ncols + 1):
            key = POSITION_FIELDNAMES[col - 1]
            val = row.get(key, "")
            c = ws.cell(row=r, column=col, value=val)
            c.alignment = cell_align

    ws.freeze_panes = "A2"

    for c in range(1, ncols + 1):
        col_letter = get_column_letter(c)
        maxlen = 0
        for r in range(1, ws.max_row + 1):
            v = ws.cell(row=r, column=c).value
            if v is not None:
                maxlen = max(maxlen, min(len(str(v)), 500))
        ws.column_dimensions[col_letter].width = min(max(maxlen + 2, 10), 55)

    wb.save(xlsx_path)
    return os.path.abspath(xlsx_path)


def process_new_markdown(
    md_path: str,
    config: Optional[dict] = None,
    dedup_index: Optional[PositionDedupIndex] = None,
) -> None:
    """
    读取 Markdown，调用 LLM 抽取「一文多岗」JSON 数组；每条有效岗位追加一行 CSV。
    去重：篇内重复 + dedup_index（CSV 已有 + 本次运行已写入）。
    """
    if config is None:
        config = load_config()

    md_path = os.path.abspath(md_path)
    if not os.path.isfile(md_path):
        print(f"  [structure_data] 文件不存在: {md_path}")
        return

    with open(md_path, "r", encoding="utf-8") as f:
        md_text = f.read()

    article_title = _resolve_article_title(md_text, md_path)
    skip_reason = _title_skip_reason(article_title)
    if skip_reason:
        _log_skip_article(md_path, skip_reason, article_title)
        return

    article_link = _article_link_from_md(md_text)
    csv_path = paths.csv_path()
    os.makedirs(os.path.dirname(csv_path) or ".", exist_ok=True)

    if dedup_index is None:
        dedup_index = PositionDedupIndex()
        loaded = dedup_index.load_csv(csv_path)
        if loaded:
            print(f"  [structure_data] 已加载 CSV 去重索引 {loaded} 条")

    user = (
        f"文章标题：{article_title}\n\n"
        f"以下为文章 Markdown：\n\n{md_text[:28000]}"
    )

    content = _llm_chat(
        [{"role": "system", "content": LLM_SYSTEM_PROMPT}, {"role": "user", "content": user}],
        config,
    )
    if content is None:
        return

    items = _extract_positions_list(content)
    if items is None:
        print("  [structure_data] 无法解析 LLM 返回的 JSON，跳过")
        return

    if len(items) == 0:
        msg = f"非招聘信息(空岗位列表): {md_path}\n"
        print(f"  [structure_data] {msg.strip()}")
        log_path = paths.skip_log_path()
        os.makedirs(os.path.dirname(log_path) or ".", exist_ok=True)
        with open(log_path, "a", encoding="utf-8") as logf:
            logf.write(msg)
        return

    items, skipped_intra = _deduplicate_llm_items(items, article_link, md_path)
    if skipped_intra:
        print(f"  [structure_data] 篇内去重跳过 {skipped_intra} 条重复岗位")

    written = 0
    skipped_dup = 0
    skipped_invalid = 0
    for zh_item in items:
        flat_preview = _flatten_zh_record_to_camel(zh_item)
        company_name = _normalize_field(flat_preview.get("companyName"))
        job_title = _normalize_field(flat_preview.get("positionName"))
        if not company_name or not job_title:
            skipped_invalid += 1
            continue
        if dedup_index.contains(article_link, company_name, job_title, md_path):
            skipped_dup += 1
            continue
        row = _zh_item_to_job_row(zh_item, article_link, md_path, article_title)
        if not _append_csv_row(row, csv_path):
            print(f"  [structure_data] CSV 写入失败（表头不匹配或 IO），中止: {md_path}")
            return
        dedup_index.add_row(row)
        written += 1

    if written == 0 and skipped_invalid == len(items):
        msg = f"非招聘信息(无有效单位/岗位): {md_path}\n"
        print(f"  [structure_data] {msg.strip()}")
        log_path = paths.skip_log_path()
        os.makedirs(os.path.dirname(log_path) or ".", exist_ok=True)
        with open(log_path, "a", encoding="utf-8") as logf:
            logf.write(msg)
    elif written == 0 and skipped_dup > 0 and skipped_invalid == 0:
        print(
            f"  [structure_data] 全部岗位已存在（去重跳过 {skipped_dup} 条）: {md_path}"
        )
    elif written == 0:
        print(
            f"  [structure_data] 未写入任何行（无效 {skipped_invalid}，去重 {skipped_dup}）: {md_path}"
        )
    else:
        dup_note = f"，去重跳过 {skipped_dup}" if skipped_dup else ""
        print(f"  [structure_data] 已写入 {written} 条岗位{dup_note}: {md_path}")


def run_batch_dir(dir_path: str, config: Optional[dict] = None) -> None:
    """递归处理目录下所有 .md；共享去重索引（CSV 已有 + 本次批量已写入）。"""
    if config is None:
        config = load_config()
    dir_path = os.path.abspath(dir_path)
    if not os.path.isdir(dir_path):
        print(f"[structure_data] 目录不存在: {dir_path}")
        return
    csv_path = paths.csv_path()
    os.makedirs(os.path.dirname(csv_path) or ".", exist_ok=True)
    dedup_index = PositionDedupIndex()
    loaded = dedup_index.load_csv(csv_path)
    print(f"[structure_data] 批量去重索引已加载 {loaded} 条（来自 {csv_path}）")
    md_files = []
    for root, _, files in os.walk(dir_path):
        for name in files:
            if name.lower().endswith(".md"):
                md_files.append(os.path.join(root, name))
    md_files.sort()
    total = len(md_files)
    print(f"[structure_data] 批量模式: {dir_path} 共 {total} 个 Markdown")
    for i, p in enumerate(md_files, 1):
        print(f"[structure_data] [{i}/{total}] {p}")
        process_new_markdown(p, config, dedup_index=dedup_index)


def main():
    paths.ensure_project_dirs()
    paths.configure(load_config())
    argv = sys.argv[1:]
    if not argv:
        print("用法:")
        print("  python structure_data.py <文章.md> [更多.md ...]     # LLM 解析指定文件")
        print("  python structure_data.py --batch-dir [目录]         # 递归处理目录下全部 .md；省略目录时默认「公众号文章」")
        print("  python structure_data.py --dedup-csv [csv路径]        # 对已有 CSV 全量去重（保留首行）")
        print("  python structure_data.py --export-xlsx [输出.xlsx]  # 从 all_positions.csv 导出 Excel")
        sys.exit(1)

    if argv[0] == "--dedup-csv":
        path = os.path.abspath(argv[1]) if len(argv) > 1 else paths.csv_path()
        deduplicate_csv_file(path)
        return

    if argv[0] == "--export-xlsx":
        out = argv[1] if len(argv) > 1 else None
        xlsx_out = None
        if out:
            xlsx_out = os.path.abspath(out)
            parent = os.path.dirname(xlsx_out)
            if parent:
                os.makedirs(parent, exist_ok=True)
        path = export_positions_xlsx(xlsx_path=xlsx_out)
        print(f"已导出: {path}")
        return

    if argv[0] == "--batch-dir":
        target = os.path.abspath(argv[1]) if len(argv) > 1 else paths.resolve_articles_dir()
        run_batch_dir(target, load_config())
        return

    cfg = load_config()
    csv_path = paths.csv_path()
    dedup_index = PositionDedupIndex()
    loaded = dedup_index.load_csv(csv_path)
    if loaded:
        print(f"[structure_data] 去重索引已加载 {loaded} 条")
    for p in argv:
        process_new_markdown(p, cfg, dedup_index=dedup_index)


if __name__ == "__main__":
    main()
